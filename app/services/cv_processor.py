# app/services/cv_processor.py
import mimetypes
import hashlib
import logging
from pathlib import Path
from typing import Optional, Dict, Any
import tempfile
import os
import datetime

import docx
import pdfplumber
import openpyxl
from langdetect import detect

from app.utils.extractors import (
    detectar_email,
    detectar_telefono,
    detectar_nombre,
    detectar_ubicacion,
    detectar_linkedin,
    detectar_github,
    extraer_experiencia_laboral,
    extraer_educacion,
    extraer_certificaciones,
    extraer_proyectos,
    extraer_idiomas,
    extraer_habilidades_avanzadas,
    detectar_experiencia,
    detectar_educacion,
    extraer_habilidades,
)
from app.integrations.supabase import supabase_client


# Configurar logs
logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s')

SUPPORTED_MIME_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}


def obtener_mime_type(file_path: Path) -> Optional[str]:
    mime_type, _ = mimetypes.guess_type(file_path)
    return mime_type


def es_duplicado(texto: str, textos_hash: set) -> bool:
    contenido_hash = hashlib.md5(texto.encode()).hexdigest()
    return contenido_hash in textos_hash


def detectar_idioma(texto: str) -> str:
    try:
        return detect(texto)
    except Exception as e:
        logger.warning(f"Fallo en detección de idioma: {e}")
        return "unknown"


def extraer_texto(file_path: Path) -> str:
    mime_type = obtener_mime_type(file_path)

    if mime_type == "application/pdf":
        with pdfplumber.open(file_path) as pdf:
            texto = " ".join([page.extract_text() or "" for page in pdf.pages])
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(file_path)
        texto = " ".join([p.text for p in doc.paragraphs])
    elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        wb = openpyxl.load_workbook(file_path)
        texto = " ".join(str(cell.value) for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row if cell.value)
    else:
        raise ValueError("Tipo de archivo no soportado")

    if not texto.strip():
        raise ValueError("Documento sin texto legible")

    return texto


def procesar_cv_avanzado(texto: str) -> Dict[str, Any]:
    """
    Procesa el texto del CV y extrae información detallada usando las nuevas funciones
    """
    try:
        # Información básica
        data_extraida = {
            "nombre": detectar_nombre(texto),
            "email": detectar_email(texto),
            "telefono": detectar_telefono(texto),
            "ubicacion": detectar_ubicacion(texto),
            "linkedin": detectar_linkedin(texto),
            "github": detectar_github(texto),
            "idioma_detectado": detectar_idioma(texto),
        }
        
        # Información detallada
        data_extraida["experiencia_laboral"] = extraer_experiencia_laboral(texto)
        data_extraida["educacion"] = extraer_educacion(texto)
        data_extraida["certificaciones"] = extraer_certificaciones(texto)
        data_extraida["proyectos"] = extraer_proyectos(texto)
        data_extraida["idiomas"] = extraer_idiomas(texto)
        data_extraida["habilidades_detalladas"] = extraer_habilidades_avanzadas(texto)
        
        # Información básica de secciones
        data_extraida["tiene_experiencia"] = detectar_experiencia(texto)
        data_extraida["tiene_educacion"] = detectar_educacion(texto)
        
        # Habilidades simples (para compatibilidad)
        habilidades_resultado = extraer_habilidades(texto)
        if isinstance(habilidades_resultado, dict):
            data_extraida["habilidades_indexadas"] = habilidades_resultado.get('habilidades', [])
            data_extraida["otras_habilidades"] = habilidades_resultado.get('otras_habilidades', [])
        else:
            data_extraida["habilidades_indexadas"] = habilidades_resultado
            data_extraida["otras_habilidades"] = []
        
        # Calcular años de experiencia aproximados
        anos_experiencia = 0
        if data_extraida["experiencia_laboral"]:
            # Contar años basado en fechas encontradas
            for exp in data_extraida["experiencia_laboral"]:
                if exp.get("fechas") and len(exp["fechas"]) >= 2:
                    try:
                        # Extraer años de las fechas
                        fechas = exp["fechas"]
                        if len(fechas) >= 2:
                            # Asumir que las fechas están en formato "Mes Año"
                            ano_inicio = fechas[0].split()[-1] if len(fechas[0].split()) > 1 else fechas[0]
                            ano_fin = fechas[1].split()[-1] if len(fechas[1].split()) > 1 else fechas[1]
                            if ano_inicio.isdigit() and ano_fin.isdigit():
                                anos_experiencia += int(ano_fin) - int(ano_inicio)
                    except:
                        pass
        
        data_extraida["experiencia_anos"] = max(anos_experiencia, 0)
        
        # Calcular puntuación de aptitud basada en factores
        puntuacion = calcular_puntuacion_aptitud(data_extraida)
        data_extraida["puntuacion_aptitud"] = puntuacion
        
        return data_extraida
        
    except Exception as e:
        logger.error(f"Error procesando CV: {e}")
        return {
            "error": str(e),
            "nombre": "Error en procesamiento",
            "email": "No encontrado",
            "telefono": "No encontrado",
            "experiencia_laboral": [],
            "educacion": [],
            "certificaciones": [],
            "proyectos": [],
            "idiomas": [],
            "habilidades_detalladas": {},
            "habilidades_indexadas": [],
            "otras_habilidades": [],
            "experiencia_anos": 0,
            "puntuacion_aptitud": 0
        }


def calcular_puntuacion_aptitud(data: Dict[str, Any]) -> int:
    """
    Calcula una puntuación de aptitud basada en la información extraída
    """
    puntuacion = 50  # Puntuación base
    
    # Experiencia laboral
    if data.get("experiencia_laboral"):
        puntuacion += min(len(data["experiencia_laboral"]) * 10, 30)
    
    # Años de experiencia
    anos_exp = data.get("experiencia_anos", 0)
    if anos_exp > 0:
        puntuacion += min(anos_exp * 2, 20)
    
    # Certificaciones
    if data.get("certificaciones"):
        puntuacion += min(len(data["certificaciones"]) * 5, 15)
    
    # Proyectos
    if data.get("proyectos"):
        puntuacion += min(len(data["proyectos"]) * 3, 10)
    
    # Habilidades técnicas
    habilidades = data.get("habilidades_detalladas", {})
    if habilidades:
        total_habilidades = sum(len(h) for h in habilidades.values())
        puntuacion += min(total_habilidades * 2, 15)
    
    # Idiomas
    if data.get("idiomas"):
        puntuacion += min(len(data["idiomas"]) * 3, 10)
    
    return min(puntuacion, 100)  # Máximo 100


async def obtener_uploader_id(documento_id: int):
    # Obtener uploader_id desde la tabla documentos_cargados
    res = supabase_client.table('documentos_cargados').select('uploader_id').eq('documento_id', documento_id).execute()
    if res.data and len(res.data) > 0:
        return res.data[0]['uploader_id']
    return None


def generar_recomendaciones(data_extraida):
    recomendaciones = []
    if data_extraida.get('habilidades_indexadas'):
        recomendaciones.append("Potencia tus habilidades en: " + ", ".join(data_extraida['habilidades_indexadas']))
    if data_extraida.get('certificaciones'):
        recomendaciones.append("Agrega más certificaciones técnicas para destacar.")
    if data_extraida.get('experiencia_anos', 0) < 3:
        recomendaciones.append("Busca proyectos freelance para ganar experiencia.")
    if data_extraida.get('puntuacion_aptitud', 0) < 70:
        recomendaciones.append("Mejora tu perfil con cursos online y proyectos personales.")
    if not recomendaciones:
        recomendaciones.append("¡Tu perfil ya es muy competitivo! Sigue así.")
    return recomendaciones


async def guardar_recomendaciones(usuario_id, recomendaciones, puntuacion_aptitud):
    for rec in recomendaciones:
        try:
            supabase_client.table('recomendaciones').insert({
                "usuario_id": usuario_id,
                "texto": rec,
                "tipo": "perfil",
                "titulo": rec[:60] if len(rec) > 0 else "Recomendación",
                "descripcion": rec,
                "prioridad": "media",
                "fecha_generacion": datetime.datetime.utcnow().isoformat(),
                "puntuacion_aptitud": puntuacion_aptitud
            }).execute()
        except Exception as e:
            print(f"Error al guardar recomendación: {e}")


async def procesar_cv(file_content: bytes, documento_id: int) -> Dict[str, Any]:
    """
    Procesa un archivo CV y extrae información relevante.
    
    Args:
        file_content: Contenido del archivo en bytes
        documento_id: ID del documento en la base de datos
        
    Returns:
        Dict con la información extraída del CV
    """
    try:
        print(f"Iniciando procesamiento de CV para documento_id: {documento_id}")
        
        # Crear archivo temporal
        with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as temp_file:
            temp_file.write(file_content)
            temp_path = temp_file.name
            print(f"Archivo temporal creado en: {temp_path}")

        try:
            # Extraer texto del archivo
            texto = extraer_texto(Path(temp_path))
            print(f"Texto extraído, longitud: {len(texto)} caracteres")
            
            # Procesar el CV con las nuevas funciones
            data_extraida = procesar_cv_avanzado(texto)
            
            # Crear resultado
            resultado = {
                "documento_id": documento_id,
                "estado": "procesado",
                "data_extraida": data_extraida,
                "texto_original": texto[:1000] + "..." if len(texto) > 1000 else texto,  # Guardar parte del texto original
                "fecha_procesamiento": datetime.datetime.utcnow().isoformat()
            }

            print("Procesamiento completado exitosamente")
            print(f"Datos extraídos: {len(data_extraida)} campos")

            # Obtener uploader_id y guardar recomendaciones
            uploader_id = await obtener_uploader_id(documento_id)
            if uploader_id:
                recomendaciones = generar_recomendaciones(data_extraida)
                puntuacion = data_extraida.get("puntuacion_aptitud")
                if puntuacion is None:
                    puntuacion = 0
                print("Puntuación a guardar en recomendaciones:", puntuacion)
                await guardar_recomendaciones(uploader_id, recomendaciones, puntuacion)

            return resultado

        finally:
            # Limpiar archivo temporal
            print(f"Eliminando archivo temporal: {temp_path}")
            os.unlink(temp_path)

    except Exception as e:
        print(f"Error en procesar_cv: {str(e)}")
        raise e